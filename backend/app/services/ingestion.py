"""Excel / CSV ingestion of officials.

Pipeline: parse -> normalise -> resolve (dedupe + entity match) -> stage.
A human reviews the staged rows (with a worked sample) and then commits.
Nothing is written to the officials table until commit.
"""

from __future__ import annotations

import csv
import io

from sqlalchemy import func, select
from sqlalchemy.orm import Session

from ..models.data_import import DataImport, ImportRow, ImportStatus, RowAction
from ..models.official import (
    Official,
    OfficialFieldProvenance,
    PROVENANCED_FIELDS,
)
from ..models.organization_unit import OrganizationUnit
from ..services import audit
from ..services.officials import recompute_verification

# canonical field -> accepted header names (compared lowercased, non-alnum stripped)
_ALIASES: dict[str, set[str]] = {
    "name": {"name", "officialname", "fullname", "contactname", "person"},
    "email": {"email", "emailid", "mail", "officialemail", "eemail"},
    "designation": {"designation", "title", "post", "role"},
    "department": {"department", "dept", "vertical"},
    "level": {"level", "orglevel", "grade", "band"},
    "location": {"location", "place", "city", "station", "posting"},
    "unit": {
        "unit", "office", "organisationunit", "organizationunit", "orgunit",
        "branch", "lho", "rbo", "ao", "circle",
    },
}
_FILLABLE = ("email", "designation", "department", "level", "location", "organization_unit_id")


def _key(header: str) -> str:
    return "".join(ch for ch in header.lower() if ch.isalnum())


def _canonical_headers(headers: list[str]) -> dict[str, str]:
    """Map a source header to its canonical name, first match wins per canonical."""
    mapping: dict[str, str] = {}
    taken: set[str] = set()
    for h in headers:
        k = _key(h)
        for canonical, names in _ALIASES.items():
            if canonical in taken:
                continue
            if k in names:
                mapping[h] = canonical
                taken.add(canonical)
                break
    return mapping


def parse_csv(text: str) -> list[dict]:
    reader = csv.DictReader(io.StringIO(text))
    return [dict(row) for row in reader]


def parse_xlsx(data: bytes) -> list[dict]:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    ws = wb.active
    rows_iter = ws.iter_rows(values_only=True)
    try:
        header = [str(c).strip() if c is not None else "" for c in next(rows_iter)]
    except StopIteration:
        return []
    out: list[dict] = []
    for values in rows_iter:
        if values is None or all(v is None for v in values):
            continue
        out.append(
            {header[i]: values[i] for i in range(min(len(header), len(values)))}
        )
    return out


def _clean(value) -> str | None:
    if value is None:
        return None
    text = " ".join(str(value).split()).strip()
    return text or None


def normalise(raw: dict, header_map: dict[str, str]) -> dict:
    norm: dict[str, str | None] = {c: None for c in _ALIASES}
    for source_header, value in raw.items():
        canonical = header_map.get(source_header)
        if canonical and norm.get(canonical) is None:
            norm[canonical] = _clean(value)
    return norm


def _resolve(
    db: Session, norm: dict, seen: set[tuple[str, str]]
) -> tuple[RowAction, int | None, int, str | None]:
    name = norm.get("name")
    if not name:
        return RowAction.SKIP, None, 0, "missing name"

    email = norm.get("email")
    dedupe_key = (name.lower(), (email or "").lower())
    if dedupe_key in seen:
        return RowAction.SKIP, None, 0, "duplicate of an earlier row in this file"
    seen.add(dedupe_key)

    if email:
        by_email = db.scalar(
            select(Official).where(func.lower(Official.email) == email.lower())
        )
        if by_email is not None:
            return RowAction.MERGE, by_email.id, 95, None

    name_matches = db.scalars(
        select(Official).where(func.lower(Official.name) == name.lower())
    ).all()
    if name_matches:
        level, unit_name = norm.get("level"), norm.get("unit")
        for off in name_matches:
            corroborated = bool(
                level and off.level and off.level.lower() == level.lower()
            )
            if not corroborated and unit_name and off.organization_unit_id:
                unit = db.get(OrganizationUnit, off.organization_unit_id)
                corroborated = bool(unit and unit.name.lower() == unit_name.lower())
            if corroborated:
                return RowAction.MERGE, off.id, 85, None
        return (
            RowAction.MERGE,
            name_matches[0].id,
            55,
            "name matches but nothing else does - review before commit",
        )

    return RowAction.CREATE, None, 70, None


def stage_import(
    db: Session, *, filename: str, rows: list[dict], actor_id: int | None
) -> DataImport:
    header_map = _canonical_headers(list(rows[0].keys())) if rows else {}

    data_import = DataImport(
        filename=filename,
        status=ImportStatus.PENDING,
        row_count=len(rows),
        created_by=actor_id,
    )
    db.add(data_import)
    db.flush()

    seen: set[tuple[str, str]] = set()
    accepted = rejected = 0
    for i, raw in enumerate(rows, start=1):
        norm = normalise(raw, header_map)
        action, official_id, confidence, error = _resolve(db, norm, seen)
        if action is RowAction.SKIP:
            rejected += 1
        else:
            accepted += 1
        db.add(
            ImportRow(
                import_id=data_import.id,
                row_number=i,
                raw={k: _clean(v) for k, v in raw.items()},
                normalized=norm,
                resolved_official_id=official_id,
                action=action,
                confidence=confidence,
                error=error,
            )
        )
    data_import.accepted_count = accepted
    data_import.rejected_count = rejected
    db.flush()

    audit.record(
        db,
        action="data_import.stage",
        entity_type="data_import",
        entity_id=data_import.id,
        actor_id=actor_id,
        after={"filename": filename, "rows": len(rows), "accepted": accepted},
    )
    return data_import


def _unit_id(db: Session, unit_name: str | None) -> int | None:
    if not unit_name:
        return None
    unit = db.scalar(
        select(OrganizationUnit).where(
            func.lower(OrganizationUnit.name) == unit_name.lower()
        )
    )
    return unit.id if unit else None


def _write_provenance(
    db: Session, official: Official, field: str, source: str, confidence: int
) -> None:
    row = db.scalar(
        select(OfficialFieldProvenance).where(
            OfficialFieldProvenance.official_id == official.id,
            OfficialFieldProvenance.field == field,
        )
    )
    if row is None:
        row = OfficialFieldProvenance(official_id=official.id, field=field)
        db.add(row)
    row.source = source
    row.confidence = confidence


def commit_import(db: Session, data_import: DataImport, *, actor_id: int | None) -> DataImport:
    if data_import.status is not ImportStatus.PENDING:
        raise ValueError("Import is not pending")

    source = f"Import: {data_import.filename}"
    created = merged = 0

    for row in data_import.rows:
        norm = row.normalized
        if row.action is RowAction.SKIP:
            continue

        unit_id = _unit_id(db, norm.get("unit"))
        values = {
            "email": norm.get("email"),
            "designation": norm.get("designation"),
            "department": norm.get("department"),
            "level": norm.get("level"),
            "location": norm.get("location"),
            "organization_unit_id": unit_id,
        }

        if row.action is RowAction.CREATE:
            official = Official(name=norm["name"], **values)
            db.add(official)
            db.flush()
            created += 1
        else:
            official = db.get(Official, row.resolved_official_id)
            if official is None:
                official = Official(name=norm["name"], **values)
                db.add(official)
                db.flush()
                created += 1
            else:
                # fill only currently-empty fields; never overwrite a real value
                for key, value in values.items():
                    if value and not getattr(official, key):
                        setattr(official, key, value)
                merged += 1

        for field in PROVENANCED_FIELDS:
            if values.get(field):
                _write_provenance(db, official, field, source, row.confidence)
        db.flush()
        recompute_verification(official)
        row.resolved_official_id = official.id

    data_import.status = ImportStatus.COMMITTED
    from datetime import datetime, timezone

    data_import.committed_at = datetime.now(timezone.utc)
    db.flush()

    audit.record(
        db,
        action="data_import.commit",
        entity_type="data_import",
        entity_id=data_import.id,
        actor_id=actor_id,
        after={"created": created, "merged": merged},
    )
    return data_import
